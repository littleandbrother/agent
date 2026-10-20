#!/usr/bin/env python3
"""Document validation only. Does not import the scientific apparatus."""
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import os
import re
import unicodedata
from pypdf import PdfReader
from docx import Document
from PIL import Image, ImageOps, ImageDraw

root=Path(__file__).resolve().parent
render=root/'qa'/'final-render'
pdf=next(render.glob('*.pdf'))
reader=PdfReader(pdf)
content=json.loads((root/'proposal_content.json').read_text())
catalog=json.loads((root/'case_catalog_v0.2.json').read_text())
docx=next(root.glob('*.docx'))
doc=Document(docx)
normalize=lambda s: re.sub(r'\s+','',unicodedata.normalize('NFKC',s))
texts=[p.extract_text() for p in reader.pages]
alltext='\n'.join(texts)
page_report=[]
for i,pg in enumerate(content['pages']):
    title=pg['title']
    matching=[j+1 for j,tx in enumerate(texts) if normalize(title) in normalize(tx)]
    page_report.append({'logical_section':i+1,'title':title,'pdf_pages':matching})
expected=[]
for pg in content['pages']:
    expected.append(pg['title'])
    for block in pg['blocks']:
        if block['kind'] in ('paragraph','subheading'): expected.append(block['text'])
        elif block['kind']=='bullets': expected.extend(block['items'])
        else:
            expected.extend(block['headers'])
            expected.extend(str(v) for row in block['rows'] for v in row)
norm=normalize(alltext)
missing=[s for s in expected if normalize(s) not in norm]
checks={
 'case_ids_complete': set(x['id'] for x in catalog['cases'])=={'W','A','B','O','C-lite','R','P-D','S'},
 'case_parameter_status_explicit': all(x['parameter_status']=='proposed-task-profile-not-validated' for x in catalog['cases']),
 'case_cards_have_complete_fields': all(len(x['fields'])>=5 for x in catalog['cases']),
 'proposal_only': catalog['experiments_started'] is False,
 'all_sections_present': all(x['pdf_pages'] for x in page_report),
 'all_authored_text_in_pdf': len(missing)==0,
 'cjk_extractable': all(x in alltext for x in ['独立贡献','最小实验','止损','暗仓','正确动态依赖','温度']),
 'no_replacement_glyphs': '\ufffd' not in alltext,
 'no_tool_citation_tokens': 'turn10view' not in alltext and '\ue200' not in alltext,
 'no_sparse_extra_pages': all(len(tx)>250 for tx in texts),
 'all_pages_rendered': len(list(render.glob('page-*.png')))==len(reader.pages),
}
with ZipFile(docx) as z:
    xml=z.read('word/document.xml').decode()
    styles=z.read('word/styles.xml').decode()
    checks['no_title_rules']='w:pBdr' not in styles
    checks['table_headers_repeat']=xml.count('w:tblHeader')==len(doc.tables)
    checks['no_external_images']=all(not x.filename.startswith('word/media/') for x in z.infolist())
secret_values=[v for k,v in os.environ.items() if ('TOKEN' in k or 'SECRET' in k or 'API_KEY' in k) and len(v)>=12]
payloads=[alltext,docx.read_bytes().decode('latin1'),(root/'proposal_content.json').read_text(),(root/'case_catalog_v0.2.json').read_text()]
checks['no_environment_credentials_in_deliverables']=not any(v in payload for v in secret_values for payload in payloads)
checks['no_environment_credentials_in_docx_xml']=not any(v in xml or v in styles for v in secret_values)
files=sorted(render.glob('page-*.png'),key=lambda p:int(p.stem.split('-')[-1]))
tw,th=300,452
for group in range((len(files)+9)//10):
    subset=files[group*10:(group+1)*10]
    sheet=Image.new('RGB',(tw*5,th*2),'#e7e9eb')
    draw=ImageDraw.Draw(sheet)
    for i,file in enumerate(subset):
        thumb=ImageOps.contain(Image.open(file).convert('RGB'),(tw-12,th-28))
        x=(i%5)*tw+6;y=(i//5)*th+23
        sheet.paste(thumb,(x,y));draw.text((x,y-18),str(group*10+i+1),fill='black')
    sheet.save(root/'qa'/f'contact-{group+1}.png')
report={'artifact_only_validation':True,'pdf_pages':len(reader.pages),'tables':len(doc.tables),
        'checks':checks,'missing_authored_text':missing,'page_sections':page_report,
        'visual_review':'pending','experiments_executed':0,'new_model_requests':0,'new_native_jobs':0,
        'artifacts':{str(p.relative_to(root)):{'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in [pdf,docx,root/'case_catalog_v0.2.json',root/'proposal_content.json']}}
(root/'qa'/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
(root/'qa'/'final_text.txt').write_text(alltext)
print(json.dumps({'pdf_pages':len(reader.pages),'checks':checks,'missing_count':len(missing),'missing_preview':missing[:3]},ensure_ascii=False))
if not all(checks.values()): raise SystemExit(1)
